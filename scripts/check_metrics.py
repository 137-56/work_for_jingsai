# scripts/check_metrics.py
"""B8 验收：报告 / PPT / 作品简介 三处取数一致性。

★ 为什么必须有这个脚本：
    "三处取数完全一致"靠人眼是核不住的 —— 报告近 5 万字、PPT 几十页、简介若干段，
    任何一处手改过一个数字就会悄悄漂移。**这种事一次都扛不住**：
    评委一旦发现同一指标在两处对不上，整份材料的可信度就打折。

两种检查：
  ① 百分比溯源：把每份材料里的所有百分比数字抓出来，标出哪些在 metrics.json
     里**找不到出处** → 找不到的不一定是错的（可能是"抽样 30%"这类描述性数字），
     但**必须逐个人工确认**
  ② 关键指标呈现：核心指标在各份材料里是否出现

用法（在工程根目录 work/ 下执行）：
    .venv/Scripts/python.exe scripts/check_metrics.py
    .venv/Scripts/python.exe scripts/check_metrics.py --docs-dir ../文档
"""
import argparse
import json
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

PCT = re.compile(r"\d+(?:\.\d+)?%")

# 核心指标（报告里都该出现）
KEY_METRICS = [
    "元素溯源覆盖率", "元素溯源准确率", "文化规则检出率", "检索 Top-5 命中率",
    "版权预警检出率", "版权预警误报率",
]


def md_text(p):
    return p.read_text(encoding="utf-8", errors="ignore")


def pptx_text(p):
    """把 pptx 里所有可见文字抓出来（含表格与备注）。"""
    try:
        from pptx import Presentation
    except ImportError:
        return None
    out = []
    try:
        prs = Presentation(str(p))
        for slide in prs.slides:
            for sh in slide.shapes:
                if sh.has_text_frame:
                    out.append(sh.text_frame.text)
                if getattr(sh, "has_table", False):
                    for row in sh.table.rows:
                        for c in row.cells:
                            out.append(c.text)
            if slide.has_notes_slide:
                out.append(slide.notes_slide.notes_text_frame.text)
    except Exception as e:
        return f"[读取失败: {type(e).__name__}: {e}]"
    return "\n".join(out)


def collect_allowed(obj, acc):
    """从 metrics.json 里收集"可溯源"的百分比写法。"""
    if isinstance(obj, dict):
        for v in obj.values():
            collect_allowed(v, acc)
    elif isinstance(obj, list):
        for v in obj:
            collect_allowed(v, acc)
    elif isinstance(obj, str):
        for m in PCT.findall(obj):
            acc.add(m)
    elif isinstance(obj, (int, float)) and not isinstance(obj, bool):
        # ★ 取绝对值：增量类的值在 metrics.json 里是负数（如 -0.40），
        #   而报告里写的是"−40.0%"——正则抓出来的数字不含负号。不取绝对值会误判成"无法溯源"。
        x = abs(obj) * 100
        if 0 <= x <= 1000:
            # ★ 生成多种写法：报告里可能写 "59%"、"59.0%"、"59.00%"
            #   只生成 "59%" 会把 "59.0%" 误判成"无法溯源"（我第一版就踩了这个）。
            acc.add(f"{x:g}%")
            acc.add(f"{x:.1f}%")
            acc.add(f"{x:.2f}%")
            for r in (1, 2):
                acc.add(f"{round(x, r):g}%")


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--docs-dir", default=str(ROOT.parent / "文档"))
    args = ap.parse_args()
    docs_dir = Path(args.docs_dir)

    mj = ROOT / "docs" / "metrics.json"
    metrics = json.loads(mj.read_text(encoding="utf-8"))

    allowed = set()
    collect_allowed(metrics, allowed)
    print(f"metrics.json 可溯源百分比写法：{len(allowed)} 种")

    # ---- 收集材料 ----
    sources = []
    for name in ("技术报告初稿.md", "作品名称与简介.md"):
        p = docs_dir / name
        if p.exists():
            sources.append((name, md_text(p), "当前口径"))
    for p in sorted(docs_dir.glob("*.pptx")):
        t = pptx_text(p)
        if t is None:
            print(f"[跳过] {p.name}：未安装 python-pptx")
            continue
        tag = "当前口径" if "纹有其源" in t else "⚠️ 疑似方案A口径（待 D 阶段重做）"
        sources.append((p.name, t, tag))

    if not sources:
        print(f"[FAIL] 在 {docs_dir} 下没找到任何待检材料")
        return 1

    # ---- ① 百分比溯源 ----
    print()
    print("=" * 74)
    print("① 各材料中的百分比 —— 是否都能在 metrics.json 找到出处")
    print("=" * 74)
    total_orphan = 0
    for name, text, tag in sources:
        found = sorted(set(PCT.findall(text)), key=lambda s: float(s[:-1]))
        orphan = [x for x in found if x not in allowed]
        total_orphan += len(orphan)
        print(f"\n  【{name}】{tag}")
        print(f"    共 {len(found)} 种百分比｜可溯源 {len(found) - len(orphan)}｜**待确认 {len(orphan)}**")
        if orphan:
            print(f"    待确认：{orphan}")
            print("    ↑ 不一定是错的（可能是描述性数字），但必须逐个确认来源")

    # ---- ② 关键指标呈现 ----
    print()
    print("=" * 74)
    print("② 核心指标在各材料中的呈现")
    print("=" * 74)
    head = f"  {'指标':<22}"
    for name, _, _ in sources:
        head += f"{name[:14]:<18}"
    print(head)
    for key in KEY_METRICS:
        line = f"  {key:<22}"
        for name, text, _ in sources:
            line += f"{('✅ 有' if key in text else '❌ 缺'):<18}"
        print(line)

    print()
    print("=" * 74)
    if total_orphan == 0:
        print("[PASS] 所有百分比均可在 metrics.json 找到出处")
        return 0
    print(f"[注意] {total_orphan} 处百分比待人工确认（见上）—— 逐条判完即可，不一定是错")
    return 0


if __name__ == "__main__":
    sys.exit(main())
